# -*- coding: utf-8 -*-
"""recon.pkl → 논문 재구성 엑셀 (openpyxl)"""
import pickle, json, datetime, sys, os
from collections import Counter, defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

REPO = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..'))
d = pickle.load(open(os.path.join(REPO, 'data', '_recon.pkl'), 'rb'))
out_rows, excl, grids = d['out_rows'], d['excl'], d['grids']
GRID = json.load(open(f'{REPO}/analysis/grid_1142.json', encoding='utf-8'))
FUNCS = ['쓰기', '말하기', '평가·채점', '읽기', '리터러시', '번역', '문법', '문학·문화', '발음', '어휘', '듣기']
METHS = ['개발', '조사', '질적·사례', '실험', '문헌·리뷰', '성능평가', '코퍼스', '기타·혼합']
FUNC_LABEL = {'리터러시': '리터러시·역량'}
METH_LABEL = {'개발': '개발연구', '조사': '조사연구', '실험': '실험연구', '코퍼스': '코퍼스분석'}
GROUPS = ['영어', '한국어', '중국어']
PAPER = {'영어': 667, '한국어': 421, '중국어': 54}

HEAD_FILL = PatternFill('solid', fgColor='1F5F66'); CODE_HEAD = PatternFill('solid', fgColor='A8701A')
HEAD_FONT = Font(bold=True, color='FFFFFF', name='맑은 고딕', size=10)
CODE_FILL = PatternFill('solid', fgColor='FFF4D6'); WARN_FILL = PatternFill('solid', fgColor='FDE2E2'); OK_FILL = PatternFill('solid', fgColor='E3F1E5')
NOTE = Font(name='맑은 고딕', size=10); BOLD = Font(bold=True, name='맑은 고딕', size=10)

def header(ws, cols, widths, code_from=None, freeze='C2'):
    ws.append(cols)
    for i, _ in enumerate(cols, 1):
        c = ws.cell(row=1, column=i)
        c.fill = CODE_HEAD if (code_from and i >= code_from) else HEAD_FILL
        c.font = HEAD_FONT; c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    ws.row_dimensions[1].height = 32; ws.freeze_panes = freeze

order = {g: i for i, g in enumerate(GROUPS)}
out_rows.sort(key=lambda o: (order[o['group']], -o['row']['year'], o['arti_id']))
n_manual = sum(1 for o in out_rows if '수동' in o['basis'] and '코딩' not in o['basis'].split('수동')[0])
n_fit = sum(1 for o in out_rows if '격자보정' in o['basis'])
n_trim = sum(1 for r, j in excl if j['basis'] == '절단')
low = [o for o in out_rows if o['score'] < 0.85]

wb = Workbook()
# ── 0. 안내 ──
ws0 = wb.active; ws0.title = '0. 안내'
lines = [
    ('논문 제출본(J1_202600078) 최종 코퍼스 1,142편 재구성본 — 영어 667 · 한국어 421 · 중국어 54', BOLD),
    ('작성 %s · 원자료: KCI Open API 재수집(corpus_full_merged.json, 원본 5시트 1,644편) · 스크립트: scope2.py + reconstruct.py + build_recon_xlsx.py' % datetime.date.today().isoformat(), NOTE),
    ('', None),
    ('■ 이 파일의 성격', BOLD),
    ('논문의 최종 코딩표(1,142편)는 유실되었다. 이 파일은 논문 3.2의 포함·제외 기준(각주 2 제외 사유 6종, 각주 3 한국어군 L1+L2)을 규칙으로 옮겨 원본 1,644편에 적용하고, 논문이 보고한 언어군 편수(667/421/54)와 그림 6 격자(기능×방법, 871편)를 제약으로 삼아 재구성한 것이다.', NOTE),
    ('따라서 언어군 합계는 논문과 정확히 같고, 기능·방법 분포는 논문 격자에 근접하지만 개별 논문의 판정은 논문 원 코딩과 다를 수 있다. 판정근거 열과 3. 경계사례 시트를 저자가 검수하면 논문 원본과의 차이를 줄일 수 있다.', NOTE),
    ('', None),
    ('■ 재구성 절차', BOLD),
    ('1) 언어군·포함 판정(규칙): 제목·주제어·초록·KCI분류에서 언어 표지(중국어/한국어·국어/영어/기타 외국어), AI 표지, 교육 표지, 메타버스·VR, 에듀테크, 번역학·NLP·문학 담론 표지를 추출해 논문 3.2 기준으로 판정. 제외사유는 각주 2의 범주(제2외국어 / 메타버스VR / AI신호부재 / 모어국어경계 / 비AI에듀테크 / 번역품질) + 교육맥락없음 / 언어비특정 / 세언어군밖.', NOTE),
    ('2) 중국어군: 후보 전편(규칙 포함 93편)을 제목·주제어 기준으로 전수 수동 판정(판정근거 "수동"). 통번역 시트의 중한 번역수업 논문은 논문 중국어군의 번역 기능이 3편인 점에 비추어 미포함으로 두고 경계사례에 올렸다.', NOTE),
    ('3) 절단: 규칙 포함 편수가 논문 편수를 넘는 영어군(746→667)·한국어군(503→421)은 규칙 신뢰도 점수(AI 표지 위치, 교육 표지, 언어 표지 위치, 시트 출처, 초록 유무) 하위 행부터 제외(판정근거 "절단", 제외사유 "기타"). 절단된 행은 모두 3. 경계사례에 있다.', NOTE),
    ('4) 기능(주)·방법(주) 코딩: 제목(3점)·주제어(2점)·초록(1점) 가중 문자열 규칙으로 후보를 뽑아 최댓값을 잠정 배정한 뒤, 논문 그림 6 격자의 언어군별 칸 편수에 맞춰 근거가 있는 후보 사이에서만 재배정(판정근거 "격자보정", 비고에 이동 경로). 기능 근거가 없는 행은 기능비특정. 중국어군 54편은 전수 수동 코딩.', NOTE),
    ('5) 핵심도구·학습자맥락(한국어군 L1/L2/미명시)은 문자열 규칙 참고값.', NOTE),
    ('', None),
    ('■ 수치 요약', BOLD),
    ('최종 1,142편 = 영어 667 · 한국어 421 · 중국어 54 (논문과 일치) · 제외 %d편 (사유별 편수는 4. 대조표)' % len(excl), NOTE),
    ('판정근거: 규칙 %d · 수동(중국어군) %d · 절단(경계) %d · 격자보정 적용 %d편' % (sum(1 for o in out_rows if o['basis'].startswith('규칙') and '격자' not in o['basis']), sum(1 for o in out_rows if o['basis'].startswith('수동')), n_trim, n_fit), NOTE),
    ('기능 특정 편수: ' + ' · '.join('%s %d/%d' % (g, sum(grids[g][0].values()), sum(grids[g][1].values())) for g in GROUPS) + ' (재구성/논문)', NOTE),
    ('', None),
    ('■ 주의', BOLD),
    ('· 논문 표 5·6(피인용)은 2026-07-01 조회값이고 이 파일의 피인용은 2026-09-29 조회값이다.', NOTE),
    ('· 초록·참고문헌은 KCI 원자료이므로 외부 제출본에서는 초록 열을 제거할 것.', NOTE),
    ('· 코딩 열(주황)은 재구성값이다. 저자가 수정하면 4. 대조표의 수식은 갱신되지 않으므로(정적 값) reconstruct.py를 다시 돌리거나 피벗으로 재집계할 것.', NOTE),
]
for t, f in lines:
    ws0.append([t])
    if f: ws0.cell(row=ws0.max_row, column=1).font = f
ws0.column_dimensions['A'].width = 160
for r in ws0.iter_rows(): r[0].alignment = Alignment(wrap_text=True, vertical='top')

# ── 1. 최종코퍼스 ──
ws1 = wb.create_sheet('1. 최종코퍼스_1142')
cols = ['순번', 'artiId', '원본시트', '언어군(확정)', '학습자맥락', '언어기능(주)', '연구방법(주)', '핵심도구', '판정근거', '신뢰도', '비고',
        '제목(국문)', '제목(영문)', '저자(소속)', '저자수', '학술지', '발행기관', '연도', '월', 'KCI분류', '주제어(국문)', '주제어(영문)',
        '초록(국문)', '초록(영문)', '피인용(KCI)', '참고문헌수', 'DOI', 'KCI링크', '규칙후보_기능', '규칙후보_방법']
widths = [6, 14, 10, 10, 9, 13, 12, 13, 14, 7, 30, 48, 40, 36, 6, 22, 20, 7, 5, 16, 30, 30, 60, 60, 9, 8, 22, 12, 24, 24]
header(ws1, cols, widths, code_from=4, freeze='E2')
for i, o in enumerate(out_rows, 1):
    r = o['row']
    ws1.append([i, o['arti_id'], r['orig_group'], o['group'], o['learner'], o['func'], o['meth'], o['tool'], o['basis'], o['score'], o['notes'],
                r['title_ko'], r['title_en'], r['authors'], r['n_authors'], r['journal'], r['publisher'], r['year'], r['month'], r['kci_field'],
                r['kw_ko'], r['kw_en'], r['abstract_ko'], r['abstract_en'], r['cited'] if r['cited'] != '' else None, r['n_refs'], r['doi'], 'KCI 원문',
                '; '.join('%s%d' % (k, v) for k, v in sorted(o['fs'].items(), key=lambda x: -x[1])),
                '; '.join('%s%d' % (k, v) for k, v in sorted(o['ms'].items(), key=lambda x: -x[1]))])
    rr = i + 1
    ws1.cell(row=rr, column=28).hyperlink = r['permalink']
    for c in range(4, 12): ws1.cell(row=rr, column=c).fill = CODE_FILL
    if o['score'] < 0.85 or '절단' in o['basis']: ws1.cell(row=rr, column=10).fill = WARN_FILL
n = len(out_rows) + 1
for col, vals in [(4, '영어,한국어,중국어'), (5, 'L1,L2,미명시'),
                  (6, '쓰기,말하기,평가·채점,읽기,리터러시·역량,번역,문법,문학·문화,발음,어휘,듣기,기능비특정'),
                  (7, '개발연구,조사연구,질적·사례,실험연구,문헌·리뷰,성능평가,코퍼스분석,기타·혼합'),
                  (8, 'ChatGPT,생성형AI일반,챗봇,기계번역,음성인식·TTS,중국어LLM,BERT·NLP,기타')]:
    dv = DataValidation(type='list', formula1='"%s"' % vals, allow_blank=True); ws1.add_data_validation(dv)
    L = get_column_letter(col); dv.add('%s2:%s%d' % (L, L, n))
ws1.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(cols)), n)

# ── 2. 제외 ──
ws2 = wb.create_sheet('2. 제외_%d' % len(excl))
excl.sort(key=lambda x: (x[1]['reason'], x[0]['orig_group'], -x[0]['year']))
header(ws2, ['순번', 'artiId', '원본시트', '제외사유', '언어군(규칙)', '판정근거', '신뢰도', '비고', '제목(국문)', '제목(영문)', '학술지', '연도', 'KCI분류', '주제어(국문)', 'KCI링크'],
       [6, 14, 10, 12, 10, 8, 7, 34, 48, 40, 22, 7, 16, 30, 12], code_from=4, freeze='E2')
for i, (r, j) in enumerate(excl, 1):
    ws2.append([i, r['arti_id'], r['orig_group'], j['reason'], j.get('group', ''), j['basis'], j['score'], '; '.join(j.get('notes', [])),
                r['title_ko'], r['title_en'], r['journal'], r['year'], r['kci_field'], r['kw_ko'], 'KCI 원문'])
    ws2.cell(row=i + 1, column=15).hyperlink = r['permalink']
    for c in range(4, 9): ws2.cell(row=i + 1, column=c).fill = CODE_FILL
ws2.auto_filter.ref = 'A1:O%d' % (len(excl) + 1)

# ── 3. 경계사례 ──
ws3 = wb.create_sheet('3. 경계사례')
header(ws3, ['구분', 'artiId', '원본시트', '언어군', '현재판정', '제외사유', '신뢰도', '비고', '제목(국문)', '학술지', '연도', '주제어(국문)', '저자 판정(기입)', 'KCI링크'],
       [14, 14, 10, 9, 8, 12, 7, 40, 50, 22, 7, 30, 14, 12], code_from=13, freeze='E2')
bd = []
for r, j in excl:
    if j['basis'] == '절단': bd.append(('절단(포함 후보)', r, j))
    elif j['basis'] == '수동' and any('경계' in x for x in j.get('notes', [])): bd.append(('수동 경계', r, j))
for o in low:
    bd.append(('포함(신뢰도 낮음)', o['row'], dict(include='Y', reason='', score=o['score'], notes=[o['notes']], group=o['group'])))
bd.sort(key=lambda x: (x[0], x[2].get('group', ''), x[2]['score']))
for i, (kind, r, j) in enumerate(bd, 1):
    ws3.append([kind, r['arti_id'], r['orig_group'], j.get('group', ''), j.get('include', 'N'), j.get('reason', ''), j['score'], '; '.join(j.get('notes', [])),
                r['title_ko'], r['journal'], r['year'], r['kw_ko'], '', 'KCI 원문'])
    ws3.cell(row=i + 1, column=14).hyperlink = r['permalink']
    ws3.cell(row=i + 1, column=13).fill = CODE_FILL
ws3.auto_filter.ref = 'A1:N%d' % (len(bd) + 1)

# ── 4. 대조표 ──
ws4 = wb.create_sheet('4. 대조표')
ws4.column_dimensions['A'].width = 26
for c in 'BCDEFGHIJ': ws4.column_dimensions[c].width = 12
def put(row, bold=False, fill=None):
    ws4.append(row)
    for c in ws4[ws4.max_row]:
        c.font = BOLD if bold else NOTE
        if fill: c.fill = fill
put(['(1) 언어군 편수', '논문', '재구성', '차이'], True, HEAD_FILL)
for c in ws4[ws4.max_row]: c.font = HEAD_FONT
cg = Counter(o['group'] for o in out_rows)
for g in GROUPS: put([g, PAPER[g], cg[g], cg[g] - PAPER[g]], fill=OK_FILL if cg[g] == PAPER[g] else WARN_FILL)
put(['합계', sum(PAPER.values()), len(out_rows), len(out_rows) - sum(PAPER.values())], True)
put([])
put(['(2) 제외 사유(각주 2 대비)', '논문 각주 2', '재구성', '비고'], True, HEAD_FILL)
for c in ws4[ws4.max_row]: c.font = HEAD_FONT
ce = Counter(j['reason'] for r, j in excl)
FN2 = {'제2외국어': 61, '메타버스VR': 42, 'AI신호부재': 34, '모어국어경계': 25, '비AI에듀테크': 24, '번역품질': 3}
for k in ['제2외국어', '메타버스VR', 'AI신호부재', '모어국어경계', '비AI에듀테크', '번역품질', '교육맥락없음', '언어비특정', '세언어군밖', '기타']:
    put([k, FN2.get(k, ''), ce.get(k, 0), '절단(경계)' if k == '기타' else ('각주 2 외 범주' if k not in FN2 else '')])
put(['합계', 189, len(excl), '논문 각주 2는 189편만 사유를 밝힘(1,644−1,142=502)'], True)
put([])
put(['(3) 표 10 언어기능(주) — 기능 특정 편수', '', '', '', '', '', ''], True, HEAD_FILL)
put(['기능', '영어 논문', '영어 재구성', '한국어 논문', '한국어 재구성', '중국어 논문', '중국어 재구성'], True)
cf = {g: Counter(o['func'] for o in out_rows if o['group'] == g) for g in GROUPS}
for f in FUNCS:
    lab = FUNC_LABEL.get(f, f); row = [lab]
    for g in GROUPS:
        p = sum(GRID[g][f].values()); rc = cf[g][lab]; row += [p, rc]
    put(row)
row = ['기능비특정']
for g in GROUPS: row += [PAPER[g] - sum(sum(m.values()) for m in GRID[g].values()), cf[g]['기능비특정']]
put(row)
put([])
put(['(4) 표 2 연구방법(주) — 기능 특정 편수(격자 열 합)', '', '', '', '', '', ''], True, HEAD_FILL)
put(['방법', '영어 논문', '영어 재구성', '한국어 논문', '한국어 재구성', '중국어 논문', '중국어 재구성'], True)
cm = {g: Counter(o['meth'] for o in out_rows if o['group'] == g and o['func'] != '기능비특정') for g in GROUPS}
cm_all = {g: Counter(o['meth'] for o in out_rows if o['group'] == g) for g in GROUPS}
for m in METHS:
    lab = METH_LABEL.get(m, m); row = [lab]
    for g in GROUPS:
        p = sum(GRID[g][f][m] for f in FUNCS); row += [p, cm[g][lab]]
    put(row)
put(['(전체 1,142 기준 재구성 방법 분포)', '', '; '.join('%s %d' % (k, v) for k, v in cm_all['영어'].most_common()), '', '; '.join('%s %d' % (k, v) for k, v in cm_all['한국어'].most_common()), '', '; '.join('%s %d' % (k, v) for k, v in cm_all['중국어'].most_common())])
put([])
put(['(5) 그림 6 격자 — 칸 불일치 요약', '', '', ''], True, HEAD_FILL)
put(['언어군', '칸 불일치 수(88칸 중)', '절대차 합', '재구성 기능특정/논문'], True)
for g in GROUPS:
    cur, tgt = grids[g]
    diff = {k: cur[k] - tgt[k] for k in tgt if cur[k] != tgt[k]}
    put([g, len(diff), sum(abs(v) for v in diff.values()), '%d/%d' % (sum(cur.values()), sum(tgt.values()))])
put([])
put(['(6) 중국어군 공백 좌표(≤1편) 수', '논문', '재구성'], True, HEAD_FILL)
cur, tgt = grids['중국어']
put(['공백 칸(중국어 ≤1)', sum(1 for k in tgt if tgt[k] <= 1), sum(1 for k in tgt if cur[k] <= 1)])
put(['0편 칸', sum(1 for k in tgt if tgt[k] == 0), sum(1 for k in tgt if cur[k] == 0)])

# ── 5. 격자 ──
ws5 = wb.create_sheet('5. 격자_재구성')
ws5.column_dimensions['A'].width = 14
for g in GROUPS:
    cur, tgt = grids[g]
    ws5.append(['%s — 재구성 (괄호: 논문 그림 6, 색: 불일치)' % g]); ws5.cell(row=ws5.max_row, column=1).font = BOLD
    ws5.append(['기능＼방법'] + METHS + ['계'])
    for c in ws5[ws5.max_row]: c.font = HEAD_FONT; c.fill = HEAD_FILL
    for f in FUNCS:
        row = [FUNC_LABEL.get(f, f)]
        for m in METHS:
            v, t = cur[(f, m)], tgt[(f, m)]
            row.append(str(v) if v == t else '%d (%d)' % (v, t))
        row.append(sum(cur[(f, m)] for m in METHS))
        ws5.append(row)
        for i, m in enumerate(METHS, 2):
            if cur[(f, m)] != tgt[(f, m)]: ws5.cell(row=ws5.max_row, column=i).fill = WARN_FILL
    ws5.append(['계'] + [sum(cur[(f, m)] for f in FUNCS) for m in METHS] + [sum(cur.values())])
    ws5.append([])
for c in 'BCDEFGHIJ': ws5.column_dimensions[c].width = 11

# ── 6. 연도×언어군 ──
ws6 = wb.create_sheet('6. 연도×언어군')
header(ws6, ['연도'] + GROUPS + ['계'], [8, 12, 12, 12, 8], freeze='B2')
piv = defaultdict(Counter)
for o in out_rows: piv[o['row']['year']][o['group']] += 1
for y in sorted(piv): ws6.append([y] + [piv[y][g] for g in GROUPS] + [sum(piv[y].values())])
ws6.append(['계'] + [cg[g] for g in GROUPS] + [len(out_rows)])

# ── 7. 한국어군 L1/L2 ──
ws7 = wb.create_sheet('7. 한국어군_L1L2')
header(ws7, ['학습자맥락', '편수', '쓰기', '말하기', '평가·채점', '읽기', '리터러시·역량', '번역', '문법', '문학·문화', '발음', '어휘', '듣기', '기능비특정'], [10, 7] + [9] * 12, freeze='B2')
for L in ['L1', 'L2', '미명시']:
    ko = [o for o in out_rows if o['group'] == '한국어' and o['learner'] == L]
    cfk = Counter(o['func'] for o in ko)
    ws7.append([L, len(ko)] + [cfk[FUNC_LABEL.get(f, f)] for f in FUNCS] + [cfk['기능비특정']])

out = sys.argv[1] if len(sys.argv) > 1 else f'{REPO}/data/KCI_논문재구성_1142_20260929.xlsx'
wb.save(out)
print('saved', out, '| 최종', len(out_rows), '| 제외', len(excl), '| 경계사례', len(bd))
print('언어군', dict(cg))
for g in GROUPS:
    print(g, '기능:', {FUNC_LABEL.get(f, f): (sum(GRID[g][f].values()), cf[g][FUNC_LABEL.get(f, f)]) for f in FUNCS}, '비특정', (PAPER[g] - sum(sum(m.values()) for m in GRID[g].values()), cf[g]['기능비특정']))
    print(g, '방법:', {METH_LABEL.get(m, m): (sum(GRID[g][f][m] for f in FUNCS), cm[g][METH_LABEL.get(m, m)]) for m in METHS})
