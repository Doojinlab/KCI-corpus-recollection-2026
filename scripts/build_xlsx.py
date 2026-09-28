# -*- coding: utf-8 -*-
"""corpus_merged.json -> KCI_AI언어교육_코퍼스_재수집.xlsx (openpyxl)"""
import json, sys, datetime
from collections import Counter, defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

rows = json.load(open('data/corpus_merged.json', encoding='utf-8'))
log = json.load(open('data/supp_search_log.json', encoding='utf-8'))
Y0, Y1 = 2020, 2026
main = [r for r in rows if isinstance(r.get('year'), int) and Y0 <= r['year'] <= Y1]
out_of_range = [r for r in rows if not (isinstance(r.get('year'), int) and Y0 <= r['year'] <= Y1)]

HEAD_FILL = PatternFill('solid', fgColor='1F5F66')
CODE_HEAD = PatternFill('solid', fgColor='A8701A')
HEAD_FONT = Font(bold=True, color='FFFFFF', name='맑은 고딕', size=10)
CODE_FILL = PatternFill('solid', fgColor='FFF4D6')
NOTE_FONT = Font(name='맑은 고딕', size=10)
BOLD = Font(bold=True, name='맑은 고딕', size=10)


def header(ws, cols, widths, code_from=None):
    ws.append(cols)
    for i, _ in enumerate(cols, 1):
        cell = ws.cell(row=1, column=i)
        cell.fill = CODE_HEAD if (code_from and i >= code_from) else HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = 'D2'


def grp(r):
    return r['orig_group'] or ('(보충)' + r['hint_group'])


wb = Workbook()
# ---------- 0. 안내 ----------
ws0 = wb.active
ws0.title = '0. 안내'
c_src = Counter(r['src'] for r in main)
c_grp = Counter(grp(r) for r in main)
c_det = sum(1 for r in main if r['detail_ok'])
c_abs = sum(1 for r in main if r['abstract_ko'] or r['abstract_en'])
n_new = sum(1 for r in rows if r['src'] == '보충')
lines = [
    ('KCI Open API 재수집 코퍼스 — AI 활용 언어교육 연구(영어·한국어·중국어), 2020–2026.06', BOLD),
    ('작성일 %s · 파이프라인: Doojinlab/KCI- kci_engine(Node.js) + build_corpus.js / build_xlsx.py' % datetime.date.today().isoformat(), NOTE_FONT),
    ('', None),
    ('■ 수집 범위와 절차', BOLD),
    ('1) 원본 1차 수집본(KCI_corpus_정리.xlsx, 2026.07)의 ①중국어 ②한국어 ③영어 시트 1,470편을 artiId 기준으로 불러옴.', NOTE_FONT),
    ('2) 보충 검색: KCI articleSearch를 제목(title)·주제어(keyword) 필드로 언어 표지 × AI 표지 조합 질의(3. 검색로그 시트) 전 페이지 수집. 심사위원 2 제안 대체 검색어(HSK, 한자, 중국어 학습, TOPIK 등), 중국어 모델명(어니봇·딥시크), 중문·영문 제목 질의, L1 국어·글쓰기 질의 포함.', NOTE_FONT),
    ('3) 원본과 보충을 artiId로 병합·중복 제거. 발행연도 2020–2026만 "1. 코퍼스_후보"에 수록(범위 밖은 "4. 연도범위밖").', NOTE_FONT),
    ('4) 전 편에 articleDetail 호출 → 초록(국·영), 주제어(국·영), 저자(소속), 발행기관, KCI 분류, DOI, 피인용(KCI), 참고문헌 목록 보강. 원시 XML은 data/cache에 보존.', NOTE_FONT),
    ('5) 규칙 태그(언어군·기능·방법·AI/교육 표지)는 제목+주제어+초록의 문자열 규칙에 따른 참고값이며 판정이 아님. 논문 3.2의 포함·제외 기준에 따른 최종 판정은 주황색 코딩 열(포함여부 이후)에 저자가 기입.', NOTE_FONT),
    ('', None),
    ('■ 수치 요약', BOLD),
    ('병합 총 %d편 = 원본 1,470편 + 보충 신규 %d편 · 이 중 2020–2026: %d편 · 범위 밖: %d편' % (len(rows), n_new, len(main), len(out_of_range)), NOTE_FONT),
    ('출처별(2020–2026): ' + ', '.join('%s %d편' % (k, v) for k, v in sorted(c_src.items())), NOTE_FONT),
    ('언어군(원본 시트 또는 보충 힌트): ' + ', '.join('%s %d편' % (k, v) for k, v in sorted(c_grp.items())), NOTE_FONT),
    ('상세 조회 확보 %d/%d편 · 초록 보유 %d/%d편 · 참고문헌 블록 보유 %d편' % (c_det, len(main), c_abs, len(main), sum(1 for r in main if r['ref_available'])), NOTE_FONT),
    ('', None),
    ('■ 코딩 열 안내(논문 3.2·4.2·5.1의 범주)', BOLD),
    ('포함여부: Y/N · 제외사유: 교육맥락없음 / 언어비특정 / 세언어군밖 / AI신호부재 / 메타버스VR / 비AI에듀테크 / 번역품질 / 모어국어경계 / 기타', NOTE_FONT),
    ('언어군(확정): 영어 / 한국어 / 중국어 · 학습자맥락: L1 / L2 / 미명시 · 언어기능(주): 쓰기 / 말하기 / 평가·채점 / 읽기 / 리터러시·역량 / 번역 / 문법 / 문학·문화 / 발음 / 어휘 / 듣기 / 기능비특정', NOTE_FONT),
    ('연구방법(주): 개발연구 / 조사연구 / 질적·사례 / 실험연구 / 문헌·리뷰 / 성능평가 / 코퍼스분석 / 기타·혼합 · 핵심도구: ChatGPT / 생성형AI일반 / 챗봇 / 기계번역 / 음성인식·TTS / 중국어LLM / BERT·NLP / 기타', NOTE_FONT),
    ('', None),
    ('■ 주의', BOLD),
    ('· 피인용은 articleDetail의 citation-count(조회일 기준)이며 논문 표 5·6의 2026.07.01 조회값과 다를 수 있음.', NOTE_FONT),
    ('· 규칙 언어군이 비어 있는 행은 규칙이 판정하지 못한 것이지 범위 밖이라는 뜻이 아님(예: 중국어 원문 제목).', NOTE_FONT),
    ('· KCI 이용약관상 원자료 재배포 제한이 있을 수 있으므로 외부 제출 시 초록 열 제외본을 별도 작성할 것.', NOTE_FONT),
]
for t, f in lines:
    ws0.append([t])
    if f:
        ws0.cell(row=ws0.max_row, column=1).font = f
ws0.column_dimensions['A'].width = 150
for r in ws0.iter_rows():
    r[0].alignment = Alignment(wrap_text=True, vertical='top')

# ---------- 1. 코퍼스_후보 ----------
ws1 = wb.create_sheet('1. 코퍼스_후보')
cols = ['순번', 'artiId', '출처', '언어군(원본시트)', '언어군(보충힌트)', '언어군(규칙)', '제목(국문)', '제목(영문)', '저자(소속)', '저자수',
        '학술지', '발행기관', '연도', '월', 'KCI분류', '주제어(국문)', '주제어(영문)', '초록(국문)', '초록(영문)', '피인용(KCI)', '참고문헌수',
        'DOI', 'KCI링크', '적중질의', '규칙태그_기능', '규칙태그_방법', 'AI표지', '교육표지',
        '포함여부', '제외사유', '언어군(확정)', '학습자맥락', '언어기능(주)', '연구방법(주)', '핵심도구', '비고']
widths = [6, 14, 9, 12, 12, 12, 48, 40, 36, 6, 22, 20, 7, 5, 14, 30, 30, 60, 60, 9, 8, 22, 12, 28, 18, 18, 7, 7,
          9, 16, 11, 10, 13, 13, 14, 20]
CODE_FROM = 29
header(ws1, cols, widths, code_from=CODE_FROM)
for i, r in enumerate(main, 1):
    cited = r['cited'] if r['cited'] != '' else None
    ws1.append([i, r['arti_id'], r['src'], r['orig_group'], r['hint_group'], r['rule_group'], r['title_ko'], r['title_en'],
                r['authors'], r['n_authors'], r['journal'], r['publisher'], r['year'], r['month'], r['kci_field'],
                r['kw_ko'], r['kw_en'], r['abstract_ko'], r['abstract_en'], cited, r['n_refs'], r['doi'], 'KCI 원문',
                r['hit_queries'], r['tag_var'], r['tag_met'], r['has_ai'], r['has_edu'], '', '', '', '', '', '', '', ''])
    rr = i + 1
    ws1.cell(row=rr, column=23).hyperlink = r['permalink']
    for c in range(CODE_FROM, len(cols) + 1):
        ws1.cell(row=rr, column=c).fill = CODE_FILL
n = len(main) + 1
lists = [
    (29, 'Y,N'),
    (30, '교육맥락없음,언어비특정,세언어군밖,AI신호부재,메타버스VR,비AI에듀테크,번역품질,모어국어경계,기타'),
    (31, '영어,한국어,중국어'),
    (32, 'L1,L2,미명시'),
    (33, '쓰기,말하기,평가·채점,읽기,리터러시·역량,번역,문법,문학·문화,발음,어휘,듣기,기능비특정'),
    (34, '개발연구,조사연구,질적·사례,실험연구,문헌·리뷰,성능평가,코퍼스분석,기타·혼합'),
    (35, 'ChatGPT,생성형AI일반,챗봇,기계번역,음성인식·TTS,중국어LLM,BERT·NLP,기타'),
]
for col, vals in lists:
    dv = DataValidation(type='list', formula1='"%s"' % vals, allow_blank=True)
    ws1.add_data_validation(dv)
    L = get_column_letter(col)
    dv.add('%s2:%s%d' % (L, L, n))
ws1.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(cols)), n)

# ---------- 2. 보충_신규 ----------
ws2 = wb.create_sheet('2. 보충_신규')
new = [r for r in main if r['src'] == '보충']
cols2 = ['순번', 'artiId', '언어군(보충힌트)', '언어군(규칙)', '제목(국문)', '제목(영문)', '저자(소속)', '학술지', '연도', 'KCI분류',
         '주제어(국문)', '적중질의', 'AI표지', '교육표지', '규칙태그_기능', '규칙태그_방법', 'KCI링크', '1차판단(저자기입)', '메모']
header(ws2, cols2, [6, 14, 12, 12, 48, 40, 32, 22, 7, 14, 30, 40, 7, 7, 18, 18, 12, 14, 24], code_from=18)
for i, r in enumerate(new, 1):
    ws2.append([i, r['arti_id'], r['hint_group'], r['rule_group'], r['title_ko'], r['title_en'], r['authors'], r['journal'],
                r['year'], r['kci_field'], r['kw_ko'], r['hit_queries'], r['has_ai'], r['has_edu'], r['tag_var'], r['tag_met'],
                'KCI 원문', '', ''])
    ws2.cell(row=i + 1, column=17).hyperlink = r['permalink']
    for c in (18, 19):
        ws2.cell(row=i + 1, column=c).fill = CODE_FILL
ws2.auto_filter.ref = 'A1:S%d' % (len(new) + 1)

# ---------- 3. 검색로그 ----------
ws3 = wb.create_sheet('3. 검색로그')
header(ws3, ['필드', '질의', '언어군', 'KCI total', '수집건', '신규건(질의 순 누적 기준)', '오류'], [9, 28, 10, 10, 9, 14, 30])
for l in log:
    ws3.append([l.get('field'), l.get('q'), l.get('group'), l.get('total'), l.get('got'), l.get('fresh'), l.get('error', '')])

# ---------- 4. 연도범위밖 ----------
ws4 = wb.create_sheet('4. 연도범위밖')
header(ws4, ['artiId', '출처', '언어군', '제목(국문)', '학술지', '연도', 'KCI링크'], [14, 9, 12, 50, 22, 7, 40])
for r in out_of_range:
    ws4.append([r['arti_id'], r['src'], grp(r), r['title_ko'], r['journal'], r['year'], r['permalink']])

# ---------- 5. 참고문헌 ----------
ws5 = wb.create_sheet('5. 참고문헌')
header(ws5, ['artiId(인용논문)', '인용논문 언어군', '인용논문 연도', '순번', '참고문헌 제목', '저자', '학술지/학회', '연도', '유형', 'DOI', 'KCI artiId(내부여부)'],
       [14, 10, 8, 5, 50, 24, 26, 7, 10, 22, 16])
ids = set(r['arti_id'] for r in main)
for r in main:
    for j, ref in enumerate(r['references'], 1):
        ra = ref.get('arti_id') or ''
        ws5.append([r['arti_id'], grp(r), r['year'], j, ref.get('title'), ref.get('author'), ref.get('journal'), ref.get('year'),
                    ref.get('type'), ref.get('doi'), ra + (' (내부)' if ra and ra in ids else '')])

# ---------- 6. 연도×언어군 ----------
ws6 = wb.create_sheet('6. 연도×언어군')
groups = sorted(c_grp.keys())
header(ws6, ['연도'] + groups + ['계'], [8] + [16] * len(groups) + [8])
piv = defaultdict(Counter)
for r in main:
    piv[r['year']][grp(r)] += 1
for y in sorted(piv):
    ws6.append([y] + [piv[y][g] for g in groups] + [sum(piv[y].values())])
ws6.append(['계'] + [c_grp[g] for g in groups] + [len(main)])

out = sys.argv[1] if len(sys.argv) > 1 else 'data/KCI_AI언어교육_코퍼스_재수집.xlsx'
wb.save(out)
print('saved', out, 'main', len(main), 'new', len(new), 'refs', ws5.max_row - 1)
